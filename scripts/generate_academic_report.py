"""
VeriScope AI - Master Academic Integrated Project Report Generator
Constructs:
  1. reports/VeriScope_AI_Integrated_Project_Report.docx
  2. reports/VeriScope_AI_Integrated_Project_Report.pdf

Follows the strict 14-section academic structure:
  1. Cover Page
  2. Declaration by Student
  3. Acknowledgement
  4. Individual Role & Responsibility
  5. Index (Table of Contents)
  6. List of Figures & List of Tables
  7. Abstract
  8. Introduction
  9. Objective
  10. Feasibility Study
  11. Segments and Their Working
  12. Conclusion
  13. Future Scope
  14. References
"""

import os
import sys
import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable, Preformatted
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas

# Import content and data
from report_content import REPORT_METADATA, METRICS_DATA, TABLES, DIAGRAMS
from report_sections_text import SECTIONS_TEXT

# ---------------------------------------------------------------------------
# COLOR PALETTE CONSTANTS
# ---------------------------------------------------------------------------
NAVY_HEX = "1E3A8A"       # Primary Academic Navy
NAVY_RGB = RGBColor(30, 58, 138)
SLATE_HEX = "334155"      # Secondary Dark Slate
SLATE_RGB = RGBColor(51, 65, 85)
MUTED_HEX = "64748B"
MUTED_RGB = RGBColor(100, 116, 139)
BLACK_RGB = RGBColor(15, 23, 42)
LIGHT_BG_HEX = "F8FAFC"
BORDER_HEX = "CBD5E1"
CALLOUT_BG_HEX = "F1F5F9"
ACCENT_BLUE_HEX = "2563EB"

# ---------------------------------------------------------------------------
# DOCX XML STYLING HELPERS
# ---------------------------------------------------------------------------
def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="{'single' if top else 'none'}" w:sz="{sz}" w:space="0" w:color="{top or 'auto'}"/>
            <w:left w:val="{'single' if left else 'none'}" w:sz="{sz}" w:space="0" w:color="{left or 'auto'}"/>
            <w:bottom w:val="{'single' if bottom else 'none'}" w:sz="{sz}" w:space="0" w:color="{bottom or 'auto'}"/>
            <w:right w:val="{'single' if right else 'none'}" w:sz="{sz}" w:space="0" w:color="{right or 'auto'}"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_docx_table(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Header Row
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], NAVY_HEX)
        set_cell_borders(hdr_cells[i], top=NAVY_HEX, bottom=NAVY_HEX, left=NAVY_HEX, right=NAVY_HEX, sz="8")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=130, right=130)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.name = "Arial"

    # Data Rows
    for row_idx, row_data in enumerate(data):
        row_cells = tbl.rows[row_idx + 1].cells
        bg_col = LIGHT_BG_HEX if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            row_cells[col_idx].text = str(val)
            set_cell_background(row_cells[col_idx], bg_col)
            set_cell_borders(row_cells[col_idx], top="E2E8F0", bottom="E2E8F0", left="E2E8F0", right="E2E8F0", sz="4")
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8.0)
                run.font.color.rgb = BLACK_RGB
                run.font.name = "Arial"

    if col_widths:
        for row in tbl.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(4)
    return tbl

def add_docx_callout(doc, text, title=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, CALLOUT_BG_HEX)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{NAVY_HEX}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=140)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        r_title = p.add_run(f"{title}\n")
        r_title.bold = True
        r_title.font.size = Pt(9.0)
        r_title.font.color.rgb = NAVY_RGB
        r_title.font.name = "Arial"
    r_body = p.add_run(text)
    r_body.font.size = Pt(8.5)
    r_body.font.italic = True
    r_body.font.color.rgb = SLATE_RGB
    r_body.font.name = "Arial"
    
    p_spc = doc.add_paragraph()
    p_spc.paragraph_format.space_after = Pt(4)

def add_docx_diagram_box(doc, content, caption=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1", sz="6")
    set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(content.strip())
    r.font.name = "Courier New"
    r.font.size = Pt(7.0)
    r.font.color.rgb = BLACK_RGB

    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(8)
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = "Arial"
        r_cap.font.size = Pt(8.0)
        r_cap.font.bold = True
        r_cap.font.color.rgb = SLATE_RGB

def add_docx_heading_1(doc, text):
    h = doc.add_heading(text, level=1)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    for run in h.runs:
        run.font.name = "Arial"
        run.font.size = Pt(15.0)
        run.font.bold = True
        run.font.color.rgb = NAVY_RGB
    return h

def add_docx_heading_2(doc, text):
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    for run in h.runs:
        run.font.name = "Arial"
        run.font.size = Pt(12.0)
        run.font.bold = True
        run.font.color.rgb = SLATE_RGB
    return h

def add_docx_body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.font.color.rgb = BLACK_RGB
    return p

# ---------------------------------------------------------------------------
# BUILD WORD (.DOCX) DELIVERABLE
# ---------------------------------------------------------------------------
def build_docx_report(filepath):
    print(f"Building Microsoft Word document: {filepath}")
    doc = docx.Document()

    # Configure A4 Page Layout
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        # Running Header
        header_p = section.header.paragraphs[0]
        header_p.text = "VeriScope AI — Intelligence Beyond the Headline | Integrated Project Report"
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header_p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8.0)
            run.font.color.rgb = MUTED_RGB

        # Running Footer
        footer_p = section.footer.paragraphs[0]
        footer_p.text = "Department of Computer Science & AI | MSc in Artificial Intelligence & Data Analytics"
        footer_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in footer_p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8.0)
            run.font.color.rgb = MUTED_RGB

    # -----------------------------------------------------------------------
    # 1. COVER PAGE
    # -----------------------------------------------------------------------
    p_title_space = doc.add_paragraph()
    p_title_space.paragraph_format.space_before = Pt(24)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(REPORT_METADATA["title"].upper())
    r_title.font.name = "Arial"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY_RGB

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    r_sub = p_sub.add_run(REPORT_METADATA["subtitle"])
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = SLATE_RGB

    # Decorative Line Box
    add_docx_callout(
        doc,
        "An Integrated Capstone Project Report Submitted in Partial Fulfillment of the Requirements for the Degree of "
        f"{REPORT_METADATA['degree']}\n{REPORT_METADATA['semester']} | {REPORT_METADATA['academic_year']}",
        "ACADEMIC SUBMISSION CERTIFICATION"
    )

    # Placeholders Meta Box Table
    meta_table = doc.add_table(rows=7, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_rows = [
        ("Candidate Name:", REPORT_METADATA["student_name"]),
        ("Roll Number / Student ID:", REPORT_METADATA["roll_number"]),
        ("Academic Programme:", REPORT_METADATA["degree"]),
        ("Academic Department:", REPORT_METADATA["department"]),
        ("University / Institution:", REPORT_METADATA["institution"]),
        ("Project Supervisor:", REPORT_METADATA["supervisor"]),
        ("Date of Submission:", REPORT_METADATA["submission_date"])
    ]
    for idx, (label, val) in enumerate(meta_rows):
        c0 = meta_table.rows[idx].cells[0]
        c1 = meta_table.rows[idx].cells[1]
        c0.text = label
        c1.text = val
        set_cell_background(c0, LIGHT_BG_HEX)
        set_cell_background(c1, "FFFFFF")
        set_cell_borders(c0, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1", sz="4")
        set_cell_borders(c1, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1", sz="4")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.size = Pt(8.5)
        c0.paragraphs[0].runs[0].font.color.rgb = SLATE_RGB
        c0.paragraphs[0].runs[0].font.name = "Arial"
        c1.paragraphs[0].runs[0].font.size = Pt(8.5)
        c1.paragraphs[0].runs[0].font.color.rgb = BLACK_RGB
        c1.paragraphs[0].runs[0].font.name = "Arial"
        c0.width = Inches(2.3)
        c1.width = Inches(4.0)

    doc.add_page_break()

    # -----------------------------------------------------------------------
    # 2. DECLARATION BY STUDENT
    # -----------------------------------------------------------------------
    add_docx_heading_1(doc, SECTIONS_TEXT["declaration"]["heading"])
    for p_text in SECTIONS_TEXT["declaration"]["body"].split("\n\n"):
        add_docx_body_paragraph(doc, p_text.strip())

    # Signature and Details Block
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(30)
    p_sign.paragraph_format.line_spacing = 1.3
    d_sign = SECTIONS_TEXT["declaration"]["sign_date_place"]
    p_sign.add_run(f"{d_sign['signature_placeholder']}\n").bold = True
    p_sign.add_run(f"Candidate Name: {d_sign['candidate_name']}\n")
    p_sign.add_run(f"Roll Number: {d_sign['roll_number']}\n")
    p_sign.add_run(f"{d_sign['date']}\n")
    p_sign.add_run(f"{d_sign['place']}\n")
    for run in p_sign.runs:
        run.font.name = "Arial"
        run.font.size = Pt(9.0)

    doc.add_page_break()

    # -----------------------------------------------------------------------
    # 3. ACKNOWLEDGEMENT
    # -----------------------------------------------------------------------
    add_docx_heading_1(doc, SECTIONS_TEXT["acknowledgement"]["heading"])
    for p_text in SECTIONS_TEXT["acknowledgement"]["body"].split("\n\n"):
        add_docx_body_paragraph(doc, p_text.strip())

    doc.add_page_break()

    # -----------------------------------------------------------------------
    # 4. INDIVIDUAL ROLE & RESPONSIBILITY
    # -----------------------------------------------------------------------
    add_docx_heading_1(doc, SECTIONS_TEXT["individual_role"]["heading"])
    add_docx_body_paragraph(doc, SECTIONS_TEXT["individual_role"]["statement"])
    
    role_headers = ["Technical Responsibility Area", "Engineering & Research Contribution Details"]
    role_data = [[item["area"], item["details"]] for item in SECTIONS_TEXT["individual_role"]["responsibilities"]]
    add_docx_table(doc, role_headers, role_data, col_widths=[2.0, 4.3])

    p_rolesign = doc.add_paragraph()
    p_rolesign.paragraph_format.space_before = Pt(20)
    p_rolesign.paragraph_format.line_spacing = 1.3
    p_rolesign.add_run("Candidate Verification:\n").bold = True
    p_rolesign.add_run(f"{REPORT_METADATA['student_name']}\n(Lead ML & NLP Engineer)\n\n")
    p_rolesign.add_run("Faculty Supervisor Endorsement:\n").bold = True
    p_rolesign.add_run(f"{REPORT_METADATA['supervisor']}\n(Project Supervisor & Academic Mentor)")
    for run in p_rolesign.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8.5)

    doc.add_page_break()

    # -----------------------------------------------------------------------
    # 5. INDEX (TABLE OF CONTENTS)
    # -----------------------------------------------------------------------
    add_docx_heading_1(doc, SECTIONS_TEXT["index"]["heading"])
    idx_headers = ["Section Number and Topic Title", "Document Page"]
    idx_data = [[title, pg] for title, pg in SECTIONS_TEXT["index"]["entries"]]
    add_docx_table(doc, idx_headers, idx_data, col_widths=[5.0, 1.3])

    doc.add_page_break()

    # -----------------------------------------------------------------------
    # 6. LIST OF FIGURES & LIST OF TABLES
    # -----------------------------------------------------------------------
    add_docx_heading_1(doc, SECTIONS_TEXT["list_figures_tables"]["heading"])
    add_docx_heading_2(doc, "List of Figures")
    fig_headers = ["Figure Number", "Figure Caption & Architecture Visualization", "Page"]
    fig_data = [[num, cap, pg] for num, cap, pg in SECTIONS_TEXT["list_figures_tables"]["figures"]]
    add_docx_table(doc, fig_headers, fig_data, col_widths=[1.2, 4.2, 0.9])

    add_docx_heading_2(doc, "List of Tables")
    tbl_headers = ["Table Number", "Table Title & Scientific Content Summary", "Page"]
    tbl_data = [[num, title, pg] for num, title, pg in SECTIONS_TEXT["list_figures_tables"]["tables"]]
    add_docx_table(doc, tbl_headers, tbl_data, col_widths=[1.2, 4.2, 0.9])

    doc.add_page_break()

    # -----------------------------------------------------------------------
    # 7. ABSTRACT
    # -----------------------------------------------------------------------
    add_docx_heading_1(doc, SECTIONS_TEXT["abstract"]["heading"])
    for p_text in SECTIONS_TEXT["abstract"]["body"].split("\n\n"):
        add_docx_body_paragraph(doc, p_text.strip())

    add_docx_callout(
        doc,
        "Empirical Benchmark Highlights (Held-Out Test Set N = 6,685):\n"
        "• Linear SVM (Calibrated): Accuracy = 98.31%, Macro F1 = 0.9831, Precision = 0.9829, Recall = 0.9835\n"
        "• Logistic Regression: Accuracy = 98.37%, Macro F1 = 0.9837, Precision = 0.9835, Recall = 0.9840\n"
        "• Multinomial Naive Bayes: Accuracy = 94.45%, Macro F1 = 0.9444, Precision = 0.9443, Recall = 0.9445\n"
        "• Leakage Remediation: Zero cross-split leakage verified across 38,583 unique body clusters.",
        "EXECUTIVE RESEARCH SUMMARY"
    )

    doc.add_page_break()

    # -----------------------------------------------------------------------
    # 8. INTRODUCTION
    # -----------------------------------------------------------------------
    add_docx_heading_1(doc, SECTIONS_TEXT["introduction"]["heading"])
    for sub in SECTIONS_TEXT["introduction"]["subsections"]:
        add_docx_heading_2(doc, sub["title"])
        for p_text in sub["text"].split("\n\n"):
            add_docx_body_paragraph(doc, p_text.strip())

    doc.add_page_break()

    # -----------------------------------------------------------------------
    # 9. OBJECTIVE
    # -----------------------------------------------------------------------
    add_docx_heading_1(doc, SECTIONS_TEXT["objective"]["heading"])
    add_docx_heading_2(doc, "9.1 Primary Aim")
    add_docx_body_paragraph(doc, SECTIONS_TEXT["objective"]["primary_aim"])
    add_docx_heading_2(doc, "9.2 Specific Technical Objectives")
    for obj in SECTIONS_TEXT["objective"]["objectives_list"]:
        add_docx_body_paragraph(doc, obj)

    doc.add_page_break()

    # -----------------------------------------------------------------------
    # 10. FEASIBILITY STUDY
    # -----------------------------------------------------------------------
    add_docx_heading_1(doc, SECTIONS_TEXT["feasibility_study"]["heading"])
    for sub in SECTIONS_TEXT["feasibility_study"]["subsections"]:
        add_docx_heading_2(doc, sub["title"])
        for p_text in sub["text"].split("\n\n"):
            add_docx_body_paragraph(doc, p_text.strip())
        if "10.5" in sub["title"]:
            add_docx_table(
                doc,
                TABLES["specifications"]["headers"],
                TABLES["specifications"]["data"],
                col_widths=[1.8, 2.3, 2.2]
            )

    doc.add_page_break()

    # -----------------------------------------------------------------------
    # 11. SEGMENTS AND THEIR WORKING
    # -----------------------------------------------------------------------
    add_docx_heading_1(doc, SECTIONS_TEXT["segments_working"]["heading"])
    for sub in SECTIONS_TEXT["segments_working"]["subsections"]:
        add_docx_heading_2(doc, sub["title"])
        for p_text in sub["text"].split("\n\n"):
            add_docx_body_paragraph(doc, p_text.strip())

        # Embed Figures and Tables in matching subsections
        if "11.1" in sub["title"]:
            add_docx_diagram_box(
                doc,
                DIAGRAMS["architecture"],
                "Figure 1: VeriScope AI End-to-End System Architecture Block Diagram"
            )
        elif "11.2" in sub["title"]:
            add_docx_table(
                doc,
                TABLES["dataset_split"]["headers"],
                TABLES["dataset_split"]["data"],
                col_widths=[1.5, 1.0, 1.1, 1.3, 1.3, 1.0]
            )
            add_docx_diagram_box(
                doc,
                DIAGRAMS["leakage_flow"],
                "Figure 2: Group-Aware Stratified Splitting and Data Leakage Elimination Flowchart"
            )
        elif "11.4" in sub["title"]:
            add_docx_table(
                doc,
                TABLES["model_comparison"]["headers"],
                TABLES["model_comparison"]["data"],
                col_widths=[1.7, 0.7, 0.7, 0.7, 0.7, 0.6, 0.6, 0.6]
            )
            add_docx_table(
                doc,
                TABLES["per_class"]["headers"],
                TABLES["per_class"]["data"],
                col_widths=[1.8, 1.0, 0.8, 0.8, 0.8, 1.1]
            )
            add_docx_table(
                doc,
                TABLES["confusion_matrices"]["headers"],
                TABLES["confusion_matrices"]["data"],
                col_widths=[1.8, 0.9, 0.9, 0.9, 0.9, 0.9]
            )
            add_docx_diagram_box(
                doc,
                DIAGRAMS["confusion_matrices"],
                "Figure 3: Comparative Confusion Matrices on Held-Out ISOT Test Set (N = 6,685)"
            )
        elif "11.5" in sub["title"]:
            add_docx_table(
                doc,
                TABLES["synthetic_benchmark"]["headers"],
                TABLES["synthetic_benchmark"]["data"],
                col_widths=[1.8, 0.8, 0.8, 0.8, 0.8, 1.3]
            )
        elif "11.6" in sub["title"]:
            add_docx_diagram_box(
                doc,
                DIAGRAMS["lime_pipeline"],
                "Figure 4: LIME Local Perturbation and Explanatory Linear Surrogate Pipeline"
            )
            add_docx_diagram_box(
                doc,
                DIAGRAMS["shap_pipeline"],
                "Figure 5: SHAP Additive Feature Attribution and Log-Odds Decomposition Flow"
            )
        elif "11.7" in sub["title"]:
            add_docx_diagram_box(
                doc,
                DIAGRAMS["nli_pipeline"],
                "Figure 6: Pretrained Cross-Encoder DeBERTa-v3 Premise-Hypothesis NLI Architecture"
            )
        elif "11.8" in sub["title"]:
            add_docx_table(
                doc,
                TABLES["bandit_arms"]["headers"],
                TABLES["bandit_arms"]["data"],
                col_widths=[0.8, 1.5, 2.0, 2.0]
            )
            add_docx_diagram_box(
                doc,
                DIAGRAMS["bandit_pipeline"],
                "Figure 7: UCB1 Contextual Multi-Armed Bandit Adaptive Retrieval Strategy Loop"
            )
        elif "11.11" in sub["title"]:
            add_docx_diagram_box(
                doc,
                DIAGRAMS["streamlit_ui"],
                "Figure 8: Streamlit Multimodal User Interface Layout and Operational Workflows"
            )
        elif "11.12" in sub["title"]:
            add_docx_table(
                doc,
                TABLES["code_mapping"]["headers"],
                TABLES["code_mapping"]["data"],
                col_widths=[1.5, 1.5, 2.0, 1.3]
            )

    doc.add_page_break()

    # -----------------------------------------------------------------------
    # 12. CONCLUSION
    # -----------------------------------------------------------------------
    add_docx_heading_1(doc, SECTIONS_TEXT["conclusion"]["heading"])
    for sub in SECTIONS_TEXT["conclusion"]["subsections"]:
        add_docx_heading_2(doc, sub["title"])
        for p_text in sub["text"].split("\n\n"):
            add_docx_body_paragraph(doc, p_text.strip())

    doc.add_page_break()

    # -----------------------------------------------------------------------
    # 13. FUTURE SCOPE
    # -----------------------------------------------------------------------
    add_docx_heading_1(doc, SECTIONS_TEXT["future_scope"]["heading"])
    for sub in SECTIONS_TEXT["future_scope"]["subsections"]:
        add_docx_heading_2(doc, sub["title"])
        add_docx_body_paragraph(doc, sub["text"])

    doc.add_page_break()

    # -----------------------------------------------------------------------
    # 14. REFERENCES
    # -----------------------------------------------------------------------
    add_docx_heading_1(doc, SECTIONS_TEXT["references"]["heading"])
    for cit in SECTIONS_TEXT["references"]["citations"]:
        p_cit = doc.add_paragraph()
        p_cit.paragraph_format.space_before = Pt(2)
        p_cit.paragraph_format.space_after = Pt(4)
        p_cit.paragraph_format.line_spacing = 1.15
        r_cit = p_cit.add_run(cit)
        r_cit.font.name = "Arial"
        r_cit.font.size = Pt(8.5)
        r_cit.font.color.rgb = BLACK_RGB

    # Save DOCX
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    doc.save(filepath)
    print(f"Successfully generated DOCX: {filepath}")

# ---------------------------------------------------------------------------
# REPORTLAB NUMBERED CANVAS
# ---------------------------------------------------------------------------
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            if self._pageNumber > 1:
                self.saveState()
                # Running Header
                self.setFont('Helvetica-Bold', 7.5)
                self.setFillColor(colors.HexColor('#1E3A8A'))
                self.drawString(54, 842 - 36, "VERISCOPE AI: INTELLIGENCE BEYOND THE HEADLINE")
                self.setFont('Helvetica', 7.5)
                self.setFillColor(colors.HexColor('#64748B'))
                self.drawRightString(595 - 54, 842 - 36, "MSc Integrated Project Report")
                self.setStrokeColor(colors.HexColor('#CBD5E1'))
                self.setLineWidth(0.5)
                self.line(54, 842 - 40, 595 - 54, 842 - 40)
                
                # Running Footer
                self.line(54, 46, 595 - 54, 46)
                self.drawString(54, 32, "Department of Computer Science & AI | MSc in Artificial Intelligence & Data Analytics")
                self.drawRightString(595 - 54, 32, f"Page {self._pageNumber} of {num_pages}")
                self.restoreState()
            super().showPage()
        super().save()

# ---------------------------------------------------------------------------
# BUILD PDF DELIVERABLE VIA REPORTLAB
# ---------------------------------------------------------------------------
def build_pdf_report(filepath):
    print(f"Building companion PDF document: {filepath}")
    doc = SimpleDocTemplate(
        filepath,
        pagesize=A4,
        leftMargin=54,   # 0.75 in
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    # Custom ReportLab Typography Styles
    title_style = ParagraphStyle(
        'DocTitle', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=20, leading=24,
        textColor=colors.HexColor(f"#{NAVY_HEX}"), alignment=1, spaceAfter=8
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle', parent=styles['Normal'],
        fontName='Helvetica-Oblique', fontSize=10.5, leading=14,
        textColor=colors.HexColor(f"#{SLATE_HEX}"), alignment=1, spaceAfter=14
    )
    h1_style = ParagraphStyle(
        'DocH1', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=13, leading=17,
        textColor=colors.HexColor(f"#{NAVY_HEX}"), spaceBefore=12, spaceAfter=6, keepWithNext=True
    )
    h2_style = ParagraphStyle(
        'DocH2', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=10.5, leading=14,
        textColor=colors.HexColor(f"#{SLATE_HEX}"), spaceBefore=8, spaceAfter=4, keepWithNext=True
    )
    body_style = ParagraphStyle(
        'DocBody', parent=styles['Normal'],
        fontName='Helvetica', fontSize=8.5, leading=12.0,
        textColor=colors.HexColor("#0F172A"), spaceBefore=0, spaceAfter=5
    )
    callout_style = ParagraphStyle(
        'DocCallout', parent=styles['Normal'],
        fontName='Helvetica-Oblique', fontSize=8.0, leading=11.5,
        textColor=colors.HexColor(f"#{SLATE_HEX}")
    )
    th_style = ParagraphStyle(
        'DocTH', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=7.5, leading=10,
        textColor=colors.white, alignment=1
    )
    td_style = ParagraphStyle(
        'DocTD', parent=styles['Normal'],
        fontName='Helvetica', fontSize=7.0, leading=9.5,
        textColor=colors.HexColor("#0F172A"), alignment=0
    )
    td_center = ParagraphStyle(
        'DocTDCenter', parent=styles['Normal'],
        fontName='Helvetica', fontSize=7.0, leading=9.5,
        textColor=colors.HexColor("#0F172A"), alignment=1
    )
    pre_style = ParagraphStyle(
        'DocPre', parent=styles['Normal'],
        fontName='Courier', fontSize=6.2, leading=7.8,
        textColor=colors.HexColor("#0F172A")
    )
    caption_style = ParagraphStyle(
        'DocCaption', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=7.5, leading=10,
        textColor=colors.HexColor(f"#{SLATE_HEX}"), alignment=1, spaceBefore=3, spaceAfter=6
    )

    story = []

    def make_rl_table(headers, data, col_widths=None):
        table_data = [[Paragraph(h, th_style) for h in headers]]
        for row in data:
            row_cells = []
            for col_idx, val in enumerate(row):
                c_style = td_style if col_idx == 0 else td_center
                row_cells.append(Paragraph(str(val), c_style))
            table_data.append(row_cells)

        t = Table(table_data, colWidths=col_widths)
        t_style = [
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f"#{NAVY_HEX}")),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{BORDER_HEX}")),
        ]
        for r_idx in range(1, len(table_data)):
            if r_idx % 2 == 0:
                t_style.append(('BACKGROUND', (0, r_idx), (-1, r_idx), colors.HexColor(f"#{LIGHT_BG_HEX}")))
        t.setStyle(TableStyle(t_style))
        return t

    def make_rl_callout(text, title=None):
        content = []
        if title:
            content.append(Paragraph(f"<b>{title}</b>", ParagraphStyle('CTitle', parent=callout_style, fontName='Helvetica-Bold', textColor=colors.HexColor(f"#{NAVY_HEX}"), fontSize=8.5, leading=11)))
            content.append(Spacer(1, 2))
        content.append(Paragraph(text.replace("\n", "<br/>"), callout_style))
        t = Table([[content]], colWidths=[487])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor(f"#{CALLOUT_BG_HEX}")),
            ('LEFTPADDING', (0, 0), (-1, -1), 10),
            ('RIGHTPADDING', (0, 0), (-1, -1), 10),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('LINELEFT', (0, 0), (-1, -1), 3.0, colors.HexColor(f"#{NAVY_HEX}")),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{BORDER_HEX}")),
        ]))
        return t

    def make_rl_diagram(art, caption=None):
        cleaned_art = art.strip()
        p_art = Preformatted(cleaned_art, pre_style)
        t = Table([[p_art]], colWidths=[487])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor(f"#{LIGHT_BG_HEX}")),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{BORDER_HEX}")),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        items = [t]
        if caption:
            items.append(Paragraph(caption, caption_style))
        return items

    # 1. Cover Page
    story.append(Spacer(1, 40))
    story.append(Paragraph(REPORT_METADATA["title"].upper(), title_style))
    story.append(Paragraph(REPORT_METADATA["subtitle"], subtitle_style))
    story.append(Spacer(1, 10))
    story.append(make_rl_callout(
        f"An Integrated Capstone Project Report Submitted in Partial Fulfillment of the Requirements for the Degree of<br/>"
        f"<b>{REPORT_METADATA['degree']}</b><br/>"
        f"{REPORT_METADATA['semester']} | {REPORT_METADATA['academic_year']}",
        "ACADEMIC SUBMISSION CERTIFICATION"
    ))
    story.append(Spacer(1, 15))

    meta_rows = [
        ["Candidate Name:", REPORT_METADATA["student_name"]],
        ["Roll Number / Student ID:", REPORT_METADATA["roll_number"]],
        ["Academic Programme:", REPORT_METADATA["degree"]],
        ["Academic Department:", REPORT_METADATA["department"]],
        ["University / Institution:", REPORT_METADATA["institution"]],
        ["Project Supervisor:", REPORT_METADATA["supervisor"]],
        ["Date of Submission:", REPORT_METADATA["submission_date"]]
    ]
    t_meta = Table([[Paragraph(f"<b>{r[0]}</b>", td_style), Paragraph(r[1], td_style)] for r in meta_rows], colWidths=[160, 327])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor(f"#{LIGHT_BG_HEX}")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{BORDER_HEX}")),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(t_meta)
    story.append(PageBreak())

    # 2. Declaration
    story.append(Paragraph(SECTIONS_TEXT["declaration"]["heading"], h1_style))
    for p_text in SECTIONS_TEXT["declaration"]["body"].split("\n\n"):
        story.append(Paragraph(p_text.strip(), body_style))
    story.append(Spacer(1, 20))
    d_sign = SECTIONS_TEXT["declaration"]["sign_date_place"]
    story.append(Paragraph(f"<b>{d_sign['signature_placeholder']}</b>", body_style))
    story.append(Paragraph(f"Candidate Name: {d_sign['candidate_name']}", body_style))
    story.append(Paragraph(f"Roll Number: {d_sign['roll_number']}", body_style))
    story.append(Paragraph(f"{d_sign['date']}", body_style))
    story.append(Paragraph(f"{d_sign['place']}", body_style))
    story.append(PageBreak())

    # 3. Acknowledgement
    story.append(Paragraph(SECTIONS_TEXT["acknowledgement"]["heading"], h1_style))
    for p_text in SECTIONS_TEXT["acknowledgement"]["body"].split("\n\n"):
        story.append(Paragraph(p_text.strip(), body_style))
    story.append(PageBreak())

    # 4. Individual Role & Responsibility
    story.append(Paragraph(SECTIONS_TEXT["individual_role"]["heading"], h1_style))
    story.append(Paragraph(SECTIONS_TEXT["individual_role"]["statement"], body_style))
    story.append(Spacer(1, 6))
    role_headers = ["Technical Responsibility Area", "Engineering & Research Contribution Details"]
    role_data = [[item["area"], item["details"]] for item in SECTIONS_TEXT["individual_role"]["responsibilities"]]
    story.append(make_rl_table(role_headers, role_data, col_widths=[140, 347]))
    story.append(Spacer(1, 12))
    story.append(Paragraph(f"<b>Candidate Verification:</b> {REPORT_METADATA['student_name']} (Lead ML & NLP Engineer)", body_style))
    story.append(Paragraph(f"<b>Faculty Endorsement:</b> {REPORT_METADATA['supervisor']} (Project Supervisor)", body_style))
    story.append(PageBreak())

    # 5. Index
    story.append(Paragraph(SECTIONS_TEXT["index"]["heading"], h1_style))
    idx_headers = ["Section Number and Topic Title", "Document Page"]
    idx_data = [[title, pg] for title, pg in SECTIONS_TEXT["index"]["entries"]]
    story.append(make_rl_table(idx_headers, idx_data, col_widths=[400, 87]))
    story.append(PageBreak())

    # 6. List of Figures & Tables
    story.append(Paragraph(SECTIONS_TEXT["list_figures_tables"]["heading"], h1_style))
    story.append(Paragraph("List of Figures", h2_style))
    fig_headers = ["Figure Number", "Figure Caption & Architecture Visualization", "Page"]
    fig_data = [[num, cap, pg] for num, cap, pg in SECTIONS_TEXT["list_figures_tables"]["figures"]]
    story.append(make_rl_table(fig_headers, fig_data, col_widths=[80, 340, 67]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("List of Tables", h2_style))
    tbl_headers = ["Table Number", "Table Title & Scientific Content Summary", "Page"]
    tbl_data = [[num, title, pg] for num, title, pg in SECTIONS_TEXT["list_figures_tables"]["tables"]]
    story.append(make_rl_table(tbl_headers, tbl_data, col_widths=[80, 340, 67]))
    story.append(PageBreak())

    # 7. Abstract
    story.append(Paragraph(SECTIONS_TEXT["abstract"]["heading"], h1_style))
    for p_text in SECTIONS_TEXT["abstract"]["body"].split("\n\n"):
        story.append(Paragraph(p_text.strip(), body_style))
    story.append(Spacer(1, 6))
    story.append(make_rl_callout(
        "Empirical Benchmark Highlights (Held-Out Test Set N = 6,685):<br/>"
        "• Linear SVM (Calibrated): Accuracy = 98.31%, Macro F1 = 0.9831, Precision = 0.9829, Recall = 0.9835<br/>"
        "• Logistic Regression: Accuracy = 98.37%, Macro F1 = 0.9837, Precision = 0.9835, Recall = 0.9840<br/>"
        "• Multinomial Naive Bayes: Accuracy = 94.45%, Macro F1 = 0.9444, Precision = 0.9443, Recall = 0.9445<br/>"
        "• Leakage Remediation: Zero cross-split leakage verified across 38,583 unique body clusters.",
        "EXECUTIVE RESEARCH SUMMARY"
    ))
    story.append(PageBreak())

    # 8. Introduction
    story.append(Paragraph(SECTIONS_TEXT["introduction"]["heading"], h1_style))
    for sub in SECTIONS_TEXT["introduction"]["subsections"]:
        story.append(Paragraph(sub["title"], h2_style))
        for p_text in sub["text"].split("\n\n"):
            story.append(Paragraph(p_text.strip(), body_style))
    story.append(PageBreak())

    # 9. Objective
    story.append(Paragraph(SECTIONS_TEXT["objective"]["heading"], h1_style))
    story.append(Paragraph("9.1 Primary Aim", h2_style))
    story.append(Paragraph(SECTIONS_TEXT["objective"]["primary_aim"], body_style))
    story.append(Paragraph("9.2 Specific Technical Objectives", h2_style))
    for obj in SECTIONS_TEXT["objective"]["objectives_list"]:
        story.append(Paragraph(obj, body_style))
    story.append(PageBreak())

    # 10. Feasibility Study
    story.append(Paragraph(SECTIONS_TEXT["feasibility_study"]["heading"], h1_style))
    for sub in SECTIONS_TEXT["feasibility_study"]["subsections"]:
        story.append(Paragraph(sub["title"], h2_style))
        for p_text in sub["text"].split("\n\n"):
            story.append(Paragraph(p_text.strip(), body_style))
        if "10.5" in sub["title"]:
            story.append(Spacer(1, 4))
            story.append(make_rl_table(
                TABLES["specifications"]["headers"],
                TABLES["specifications"]["data"],
                col_widths=[120, 183, 184]
            ))
            story.append(Spacer(1, 4))
    story.append(PageBreak())

    # 11. Segments and Their Working
    story.append(Paragraph(SECTIONS_TEXT["segments_working"]["heading"], h1_style))
    for sub in SECTIONS_TEXT["segments_working"]["subsections"]:
        story.append(Paragraph(sub["title"], h2_style))
        for p_text in sub["text"].split("\n\n"):
            story.append(Paragraph(p_text.strip(), body_style))

        if "11.1" in sub["title"]:
            for item in make_rl_diagram(DIAGRAMS["architecture"], "Figure 1: VeriScope AI End-to-End System Architecture Block Diagram"):
                story.append(item)
        elif "11.2" in sub["title"]:
            story.append(make_rl_table(
                TABLES["dataset_split"]["headers"],
                TABLES["dataset_split"]["data"],
                col_widths=[110, 75, 80, 82, 80, 60]
            ))
            story.append(Spacer(1, 4))
            for item in make_rl_diagram(DIAGRAMS["leakage_flow"], "Figure 2: Group-Aware Stratified Splitting and Data Leakage Elimination Flowchart"):
                story.append(item)
        elif "11.4" in sub["title"]:
            story.append(make_rl_table(
                TABLES["model_comparison"]["headers"],
                TABLES["model_comparison"]["data"],
                col_widths=[127, 50, 52, 52, 52, 52, 52, 50]
            ))
            story.append(Spacer(1, 4))
            story.append(make_rl_table(
                TABLES["per_class"]["headers"],
                TABLES["per_class"]["data"],
                col_widths=[137, 75, 65, 65, 65, 80]
            ))
            story.append(Spacer(1, 4))
            story.append(make_rl_table(
                TABLES["confusion_matrices"]["headers"],
                TABLES["confusion_matrices"]["data"],
                col_widths=[137, 70, 70, 70, 70, 70]
            ))
            story.append(Spacer(1, 4))
            for item in make_rl_diagram(DIAGRAMS["confusion_matrices"], "Figure 3: Comparative Confusion Matrices on Held-Out ISOT Test Set (N = 6,685)"):
                story.append(item)
        elif "11.5" in sub["title"]:
            story.append(make_rl_table(
                TABLES["synthetic_benchmark"]["headers"],
                TABLES["synthetic_benchmark"]["data"],
                col_widths=[137, 65, 65, 65, 65, 90]
            ))
            story.append(Spacer(1, 4))
        elif "11.6" in sub["title"]:
            for item in make_rl_diagram(DIAGRAMS["lime_pipeline"], "Figure 4: LIME Local Perturbation and Explanatory Linear Surrogate Pipeline"):
                story.append(item)
            for item in make_rl_diagram(DIAGRAMS["shap_pipeline"], "Figure 5: SHAP Additive Feature Attribution and Log-Odds Decomposition Flow"):
                story.append(item)
        elif "11.7" in sub["title"]:
            for item in make_rl_diagram(DIAGRAMS["nli_pipeline"], "Figure 6: Pretrained Cross-Encoder DeBERTa-v3 Premise-Hypothesis NLI Architecture"):
                story.append(item)
        elif "11.8" in sub["title"]:
            story.append(make_rl_table(
                TABLES["bandit_arms"]["headers"],
                TABLES["bandit_arms"]["data"],
                col_widths=[60, 110, 157, 160]
            ))
            story.append(Spacer(1, 4))
            for item in make_rl_diagram(DIAGRAMS["bandit_pipeline"], "Figure 7: UCB1 Contextual Multi-Armed Bandit Adaptive Retrieval Strategy Loop"):
                story.append(item)
        elif "11.11" in sub["title"]:
            for item in make_rl_diagram(DIAGRAMS["streamlit_ui"], "Figure 8: Streamlit Multimodal User Interface Layout and Operational Workflows"):
                story.append(item)
        elif "11.12" in sub["title"]:
            story.append(make_rl_table(
                TABLES["code_mapping"]["headers"],
                TABLES["code_mapping"]["data"],
                col_widths=[120, 110, 147, 110]
            ))
            story.append(Spacer(1, 4))

    story.append(PageBreak())

    # 12. Conclusion
    story.append(Paragraph(SECTIONS_TEXT["conclusion"]["heading"], h1_style))
    for sub in SECTIONS_TEXT["conclusion"]["subsections"]:
        story.append(Paragraph(sub["title"], h2_style))
        for p_text in sub["text"].split("\n\n"):
            story.append(Paragraph(p_text.strip(), body_style))
    story.append(PageBreak())

    # 13. Future Scope
    story.append(Paragraph(SECTIONS_TEXT["future_scope"]["heading"], h1_style))
    for sub in SECTIONS_TEXT["future_scope"]["subsections"]:
        story.append(Paragraph(sub["title"], h2_style))
        story.append(Paragraph(sub["text"], body_style))
    story.append(PageBreak())

    # 14. References
    story.append(Paragraph(SECTIONS_TEXT["references"]["heading"], h1_style))
    for cit in SECTIONS_TEXT["references"]["citations"]:
        story.append(Paragraph(cit, body_style))
        story.append(Spacer(1, 2))

    # Build PDF with NumberedCanvas
    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"Successfully generated PDF: {filepath}")

# ---------------------------------------------------------------------------
# MAIN EXECUTION
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    report_dir = "reports"
    docx_path = os.path.join(report_dir, "VeriScope_AI_Integrated_Project_Report.docx")
    pdf_path = os.path.join(report_dir, "VeriScope_AI_Integrated_Project_Report.pdf")

    print("=" * 70)
    print("VERISCOPE AI: ACADEMIC PROJECT REPORT GENERATION")
    print("=" * 70)

    try:
        build_docx_report(docx_path)
        build_pdf_report(pdf_path)
        print("=" * 70)
        print("REPORT GENERATION COMPLETE!")
        print(f"Editable Word Document: {docx_path}")
        print(f"Companion PDF Document: {pdf_path}")
        print("=" * 70)
    except Exception as e:
        print(f"ERROR GENERATING REPORT: {e}", file=sys.stderr)
        import traceback
        traceback.print_exc()
        sys.exit(1)
